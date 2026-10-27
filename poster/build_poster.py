#!/usr/bin/env python
"""Gera poster/poster_final.pptx (a partir do template oficial do professor) e poster/poster_final.html
com o MESMO conteúdo e a MESMA geometria (cm, em poster/layout.py). Com --pdf, imprime o HTML em
poster/poster_final.pdf pelo Chrome headless (mudo), grava título e autores nos metadados do PDF,
refaz poster/poster_final_preview.png e confere as saídas.

Uso:  python poster/make_assets.py && python poster/build_poster.py --pdf
Requer: python-pptx, Pillow, pymupdf; poster/numeros.json (src/experiments/make_paper_numbers.py);
a fonte Arial do sistema (macOS: /System/Library/Fonts/Supplemental).

QUEBRA DE LINHA: o texto é quebrado AQUI, com as medidas da Arial, e cada linha vai com uma quebra
explícita para o .pptx e para o .html. Assim o PowerPoint, o Quick Look e o Chrome mostram as mesmas
linhas (antes cada programa quebrava a seu modo e o .pptx saía diferente do PDF). Cada linha ocupa no
máximo FOLGA da largura da caixa. As caixas de texto são empilhadas pela altura medida, e o script
para com erro se algo não couber.

NÚMEROS: todo número de resultado vem de poster/numeros.json (chave "valores"; mesma fonte do
paper/numeros.tex; origem de cada um em results/numeros_resumo.md). As únicas exceções são números da
LITERATURA, com a fonte comentada ao lado (27,8 % ABRAMET/PRF; > 90 % no NTHU-DDD, Lyu et al. 2018).
Aparelhos: o PP2 mostra cada aparelho medido em numeros.json (iPhone 17 Pro; o Galaxy S23 Ultra só
aparece com samsung_medido = true), com a versão do app de cada medição.
"""
import argparse
import datetime
import html
import json
import os
import re
import subprocess
import sys
import tempfile
import time

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
POSTER = os.path.join(ROOT, "poster")
ASSETS = os.path.join(POSTER, "assets")
sys.path.insert(0, POSTER)
sys.dont_write_bytecode = True   # não cria poster/__pycache__
import layout as L  # noqa: E402

TEMPLATE = os.environ.get(   # template do professor: tcc/material_disciplina/, ao lado do repositório
    "POSTER_TEMPLATE", os.path.join(os.path.dirname(os.path.dirname(POSTER)), "material_disciplina", "TemplatePoster.pptx")
)
OUT_PPTX = os.path.join(POSTER, "poster_final.pptx")
OUT_HTML = os.path.join(POSTER, "poster_final.html")
OUT_PDF = os.path.join(POSTER, "poster_final.pdf")
OUT_PNG = os.path.join(POSTER, "poster_final_preview.png")
CHROME = os.environ.get("CHROME", "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")

# Cores: cinzas e o vermelho do título do template (o semáforo do app fica só na figura do método)
RED = "D52128"
INK = "1A1A1A"
GRAY = "595959"
RULE = "8C8C8C"
FONT = "Arial"          # a mesma no .pptx e no PDF (a Montserrat do template não está instalada)
FONTES_TTF = {   # Arial normal e negrito: macOS, Windows, Linux (msttcorefonts)
    False: ["/System/Library/Fonts/Supplemental/Arial.ttf", "/Library/Fonts/Arial.ttf", "C:/Windows/Fonts/arial.ttf",
            "/usr/share/fonts/truetype/msttcorefonts/Arial.ttf"],
    True: ["/System/Library/Fonts/Supplemental/Arial Bold.ttf", "/Library/Fonts/Arial Bold.ttf",
           "C:/Windows/Fonts/arialbd.ttf", "/usr/share/fonts/truetype/msttcorefonts/Arial_Bold.ttf"]}
FOLGA = 0.95            # cada linha ocupa no máximo 95 % da largura da caixa
PT_CM = 2.54 / 72
NB = chr(0xA0)           # espaço não separável (mantém "17 ms" na mesma linha)

TITULO = ("Detecção de Sonolência em Dispositivos Móveis: Avaliação da Generalização entre "
          "Sonolência Atuada e Real")
AUTORES_COMPLETOS = ["André Luís Oliveira Ferreira da Silva", "Bruno Rodolfo da Silveira",
                     "Bruno Sandoval Ribeiro", "Paulo Henrique Perin"]
ASSUNTO = "Pôster do Projeto Transformador II, Bacharelado em Ciência da Computação, PUCPR, 2026"
PALAVRAS = "detecção de sonolência; sonolência atuada; sonolência real; UTA-RLDD; DMD; celular"
REPO_URL = "github.com/BrunoRodolfoo/deteccao-sonolencia"

SLIDE_W, SLIDE_H = L.SLIDE_W, L.SLIDE_H

# --------------------------------------------------------------------------------------
# Números (poster/numeros.json)
# --------------------------------------------------------------------------------------
with open(os.path.join(POSTER, "numeros.json"), encoding="utf-8") as _f:
    _NUM = json.load(_f)
V = _NUM["valores"]
SAMSUNG_MEDIDO = bool(_NUM.get("samsung_medido"))


def N(nome):
    """Valor pt-BR de um macro de paper/numeros.tex (erro se não existir: nada é digitado)."""
    if nome not in V:
        raise KeyError(f"número {nome!r} não está em poster/numeros.json")
    return V[nome]


def _int(nome):
    """Valor pt-BR de numeros.json arredondado para inteiro (meio para cima): '40,5' -> '41'."""
    return str(int(float(N(nome).replace(".", "").replace(",", ".")) + 0.5))


# --------------------------------------------------------------------------------------
# Parágrafos: dict(runs=[(texto, negrito, cor|None, tamanho|None)], size, sb, align, color, lh)
# --------------------------------------------------------------------------------------
def P(*runs, size=20, sb=0, align="l", color=INK, lh=1.1):
    rr = []
    for r in runs:
        if isinstance(r, str):
            rr.append((r, False, None, None))
        elif len(r) == 2:
            rr.append((r[0], r[1], None, None))
        else:
            rr.append(tuple(r) + (None,) * (4 - len(r)))
    return dict(runs=rr, size=size, sb=sb, align=align, color=color, lh=lh)


def B(t, color=None, size=None):
    return (t, True, color, size)


def C(t, color, size=None):
    return (t, False, color, size)


# --------------------------------------------------------------------------------------
# Medida do texto e quebra de linha (Arial do sistema)
# --------------------------------------------------------------------------------------
_TTF = {}


def largura_cm(texto, negrito, tamanho):
    if negrito not in _TTF:
        achou = [c for c in FONTES_TTF[negrito] if os.path.exists(c)]
        if not achou:
            raise SystemExit(f"Arial não encontrada em {FONTES_TTF[negrito]} (a quebra de linha usa a Arial)")
        _TTF[negrito] = ImageFont.truetype(achou[0], 1000)
    return _TTF[negrito].getlength(texto) * tamanho / 1000 * PT_CM


def linha_pt(spec):
    """Altura de cada linha do parágrafo (pt): a maior fonte do parágrafo x lh."""
    maior = max([spec["size"]] + [r[3] for r in spec["runs"] if r[3]])
    return round(maior * spec["lh"], 1)


def quebrar(spec, larg_cm):
    """Lista de linhas; cada linha é uma lista de runs (texto, negrito, cor, tamanho). Quebra só em
    espaços normais (o NB não quebra) e em '\\n' explícito."""
    lim = FOLGA * larg_cm
    toks, atual = [], []
    for texto, neg, cor, tam in spec["runs"]:
        cor, tam = cor or spec["color"], tam or spec["size"]
        for parte in re.split(r"( |\n)", texto):
            if not parte:
                continue
            if parte in (" ", "\n"):
                if atual:
                    toks.append(("p", atual))
                    atual = []
                toks.append(("e" if parte == " " else "n", (" ", neg, cor, tam)))
            else:
                atual.append((parte, neg, cor, tam))
    if atual:
        toks.append(("p", atual))
    linhas, linha, w, esp = [], [], 0.0, None
    for tipo, v in toks:
        if tipo == "n":
            linhas.append(linha)
            linha, w, esp = [], 0.0, None
            continue
        if tipo == "e":
            esp = v
            continue
        wp = sum(largura_cm(t, n, s) for t, n, _, s in v)
        if wp > lim:
            raise SystemExit(f"palavra maior que a caixa: {''.join(t for t, *_ in v)!r} ({wp:.2f} > {lim:.2f} cm)")
        we = largura_cm(" ", esp[1], esp[3]) if (linha and esp) else 0.0
        if linha and w + we + wp > lim:
            linhas.append(linha)
            linha, w, we = [], 0.0, 0.0
        if linha and esp:
            linha.append(esp)
            w += we
        linha.extend(v)
        w += wp
        esp = None
    if linha:
        linhas.append(linha)
    juntas = []   # junta runs vizinhos com o mesmo estilo
    for ln in linhas:
        j = []
        for t, n, c, s in ln:
            if j and j[-1][1:] == (n, c, s):
                j[-1] = (j[-1][0] + t, n, c, s)
            else:
                j.append((t, n, c, s))
        juntas.append(j)
    return juntas


def preparar(paras, larg_cm):
    """Copia os parágrafos com as linhas já quebradas para a largura dada."""
    out = []
    for sp in paras:
        sp = dict(sp)
        sp["linhas"] = quebrar(sp, larg_cm)
        sp["larg_max"] = max((sum(largura_cm(t, n, s) for t, n, _, s in ln) for ln in sp["linhas"]), default=0)
        out.append(sp)
    return out


def altura_cm(paras):
    h = 0.0
    for i, sp in enumerate(paras):
        h += (sp["sb"] if i else 0) * PT_CM + len(sp["linhas"]) * linha_pt(sp) * PT_CM
    return h


# --------------------------------------------------------------------------------------
# Conteúdo
# --------------------------------------------------------------------------------------
# Título em 3 linhas fixas (32 pt): a primeira ocupa 94,7 % da caixa
TITULO_LINHAS = "Detecção de Sonolência em Dispositivos\nMóveis: Avaliação da Generalização\nentre Sonolência Atuada e Real"
assert TITULO_LINHAS.replace("\n", " ") == TITULO
HEAD_TITLE = [P(B(TITULO_LINHAS, RED), size=32, align="c", lh=1.05)]
HEAD_URL = [P(C("Código e app: ", GRAY), C(REPO_URL, GRAY), size=16, align="c")]
AUTHORS = [
    P("André Luís O. F. da Silva"),
    P("Bruno Rodolfo da Silveira"),
    P("Bruno Sandoval Ribeiro"),
    P("Paulo Henrique Perin"),
    P(C("Orientador:", GRAY), sb=4),
    P("Prof. Júlio César Nievola"),
]

# ---- Coluna 1: Contexto e Pergunta ----------------------------------------------------------
# Literatura: 27,8 % = ABRAMET sobre dados da PRF 2014-2024 (paper/fontes_brasil.md, abramet2026);
# > 90 % = LMDF 90,05 % no NTHU-DDD (lyu2018; paper/related_work_notes.md).
COL1_TITLE = "Contexto e Pergunta"
COL1 = [
    P("Sono e saúde do condutor pesam em ", B("27,8%"), " das ocorrências nas rodovias federais [7]."),
    P("Com sono encenado (atuado), os detectores passam de 90% de acurácia [6]. Com sono real (UTA-RLDD), "
      f"o melhor modelo acerta {N('ExpUmHMLSTMVA')}% dos vídeos, e pessoas, {N('ExpUmHumanosVA')}% [1].",
      sb=14),
    P(B("Pergunta central", RED), sb=30),
    P("Um detector treinado com sono encenado reconhece sono real de pessoas nunca vistas, no celular?", sb=4),
    P(B("Perguntas de pesquisa", RED), sb=30),
    P(B("PP1."), " Quanto se perde do sono encenado para o real?", sb=4),
    P(B("PP2."), " Qual o custo no celular?", sb=10),
    P(B("PP3."), " Resiste à noite, ao movimento e aos óculos?", sb=10),
]

# ---- Coluna 2: Método e Resultados ----------------------------------------------------------
COL2_TITLE = "Método e Resultados"
TABLE = [
    ["UTA-RLDD [1]", f"sono real, {N('BaseUTASujeitos')} pessoas", "teste (PP1)"],
    ["DMD [2]", f"encenado, {N('BaseDMDSujeitos')} pessoas", "treino (PP1)"],
    ["NITYMED [3]", "noite, carro andando", "noite (PP3)"],
    ["MRL + CEW [4, 5]", "recortes de olho", "olho e óculos (PP3)"],
]
# Exp. 5b (figura do PP1): B = treino e teste na UTA-RLDD; C = treino no DMD e teste na UTA-RLDD; a perda
# mostrada é B − C. Os rótulos ficam no próprio gráfico (make_assets.py), sem legenda acima.
# Vazamento (Exp. 2 e Exp. 5b): sem número em destaque; não explica os 90 % da literatura, que foram
# medidos com outras pessoas no teste
LEAK = [
    P(B("Vazamento")),
    P("Nas sequências de piscadas, que se sobrepõem, partir por janela infla a acurácia em "
      f"{N('ExpDoisRFInflacaoJanelaBSA')}{NB}pp; nas janelas de 10 s, {N('ExpCincoBVazamentoAcBal')}{NB}pp "
      f"no RF e {N('ExpCincoBLRVazamentoAcBal')}{NB}pp na RL."),
]
# App v2.1: teste independente em sono real (UTA-RLDD, pessoas que não entraram no ajuste das regras).
# Substitui a frase da direção real do DMD (3,3 % em vermelho), que usava os mesmos vídeos do ajuste.
APP = [
    P(B(f"App v{N('AppVersao')}")),
    P(f"Em {N('AppUTASujeitos')} pessoas com sono real, fora do ajuste, separou sonolento de acordado "
      f"(AUC{NB}{N('AppUTAAUCForaVerde')}); alarmou {N('AppUTASonolentosVermelho')} sonolentos e "
      f"{N('AppUTAAcordadosAlarme')} acordados."),
]
APP_RESERVA_LINHAS = 0   # linhas guardadas na caixa do app para um resultado futuro

# ---- Coluna 3: Conclusões -------------------------------------------------------------------
COL3_TITLE = "Conclusões"
SINTESE = [P(f"Sono encenado é mais extremo que o real: o modelo do DMD reconhece só "
             f"{_int('ExpCincoBCSensJanela')}% das janelas de sono real.", size=24)]


def _faixa_ms(pre):
    """'11 a 23 ms' (faixa do p50 entre as execuções) ou '17 ms' se houver um valor só."""
    lo, hi = _int(f"{pre}TotalPcinquentaMin"), _int(f"{pre}TotalPcinquentaMax")
    return f"{lo}{NB}ms" if lo == hi else f"{lo} a {hi}{NB}ms"


def _aparelhos():
    """PP2 com cada aparelho medido em numeros.json e a versão do app de cada medição."""
    aps = [("iPhone 17 Pro", "iPhone", "CelularIPhone")]
    if SAMSUNG_MEDIDO:   # após bash tools/atualizar_celular.sh com o benchmark do Galaxy S23 Ultra
        aps.append(("Galaxy S23 Ultra", "Galaxy", "CelularSamsung"))
    vers = [N(pre + "VersaoApp") for _, _, pre in aps]
    if len(aps) == 1:
        return f"No {aps[0][0]}, {_faixa_ms(aps[0][2])} por quadro (app {vers[0]})."
    partes = " e ".join(f"{_faixa_ms(pre)} no {nome}" for nome, _, pre in aps)
    app = (f"app {vers[0]}" if len(set(vers)) == 1 else
           "app " + " e ".join(f"{v} no {curto}" for (_, curto, _), v in zip(aps, vers)))
    return f"Por quadro, {partes} ({app})."


COL3 = [
    P(B("PP1."), f" A RL do DMD perde de {N('ExpCincoBLRQuedaNormMin')} a {N('ExpCincoBLRQuedaNormMax')} pontos "
      f"de AUC em todas as normalizações testadas; o RF perde {N('ExpCincoBQuedaAUCJanela')} só com a referência "
      "principal na UTA."),
    P(B("PP2."), " " + _aparelhos(), sb=14),
    P(B("PP3."), f" A SNR do EAR (abertura do olho) é {N('ExpSeisSNRNoite')} à noite e "
      f"{N('ExpSeisSNRDMDMovimento')} com o carro andando de dia. Óculos: AUC do olho de "
      f"{N('ExpTresCEWPatchMRLSemOculosAUC')}{NB}para{NB}{N('ExpTresCEWPatchMRLComOculosAUC')}.", sb=14),
    P(B("Limitações."), " Um só par de bases.", sb=14),
]
# Referências (16 pt), em texto corrido: chaves de paper/refs.bib (NITYMED: dataset + artigo dos autores)
_REFS = ["Ghoddoosian et al., CVPRW 2019.", "Ortega et al., ECCVW 2020.", "Petrellis et al., IEEE DataPort 2022.",
         "Fusek, ISVC 2018.", "Song et al., Pattern Recognit. 2014.", "Lyu et al., arXiv 2018.",
         "ABRAMET/PRF, Agência Brasil 2026."]
REFS = [P(" ".join(f"[{i}]{NB}{r}" for i, r in enumerate(_REFS, 1)), size=16, color=GRAY, lh=1.02)]

# ======================================================================================
# Elementos: a MESMA lista gera o .pptx e o .html
#   ("card", x, y, w, h, None, {})   caixa branca do template
#   ("text", x, y, w, h, paras, kw)  texto (paras já quebrados em linhas)
#   ("img",  x, y, w, h, nome, {})   figura de poster/assets
#   ("rule", x, y, w, h, cor, {})    retângulo fino (linhas da tabela)
# ======================================================================================
def pic_h(name, w):
    pw, ph = Image.open(os.path.join(ASSETS, name)).size
    return round(w * ph / pw, 3)


def elements():
    el = [("card", x, y, w, h, None, {}) for (x, y, w, h) in L.CARDS]

    def T(x, y, w, paras, name, h=None, **kw):
        pp = preparar(paras, w)
        hh = altura_cm(pp)
        if h is not None and hh > h + 1e-6 and kw.pop("ajustavel", False):
            # se não couber, reduz só o espaço entre linhas e entre parágrafos (nunca a fonte), até 1,02
            for lh in (1.09, 1.08, 1.07, 1.06, 1.05, 1.04, 1.03, 1.02):
                fator = (lh - 1) / 0.1
                pp = preparar([dict(sp, lh=lh, sb=sp["sb"] * fator) for sp in paras], w)
                hh = altura_cm(pp)
                if hh <= h + 1e-6:
                    print(f"{name}: espaço entre linhas reduzido para {lh:.2f} para caber")
                    break
        kw.pop("ajustavel", None)
        if h is not None and hh > h + 1e-6:
            raise SystemExit(f"{name}: o texto ({hh:.2f} cm) não cabe na altura {h:.2f} cm")
        el.append(("text", x, y, w, h if h is not None else round(hh + 0.05, 3), pp, dict(name=name, **kw)))
        return hh

    def I(name, x, y, w, h=None):
        if name.startswith("fig_"):   # figuras com texto em pt reais: entram no tamanho em que foram geradas
            pw = Image.open(os.path.join(ASSETS, name)).size[0] / 300 * 2.54
            if abs(pw - w) > 0.05:
                raise SystemExit(f"{name} tem {pw:.2f} cm de largura e entraria com {w:.2f} cm: rode make_assets.py")
        el.append(("img", x, y, w, h if h else pic_h(name, w), name, {}))

    # cabeçalho: título e endereço centrados na caixa do meio; autores e QR na caixa da direita
    cab_y, cab_h = L.CARDS[1][1], L.CARDS[1][3]
    ht = altura_cm(preparar(HEAD_TITLE, L.TITLE_W))
    hu = altura_cm(preparar(HEAD_URL, L.TITLE_W))
    y0 = cab_y + (cab_h - (ht + 0.3 + hu)) / 2
    T(L.TITLE_X, y0, L.TITLE_W, HEAD_TITLE, "Titulo", cls="title")
    T(L.TITLE_X, y0 + ht + 0.3, L.TITLE_W, HEAD_URL, "Endereco do repositorio", cls="ref")
    ha = altura_cm(preparar(AUTHORS, L.AUT_W))
    aut_y, aut_h = L.CARDS[2][1], L.CARDS[2][3]
    T(L.AUT_X, aut_y + (aut_h - ha) / 2, L.AUT_W, AUTHORS, "Autores")
    I("qr_repositorio.png", L.QR_X, aut_y + (aut_h - L.QR_SIDE) / 2, L.QR_SIDE, L.QR_SIDE)

    # coluna 1: título (pode ter 2 linhas) e texto
    h1 = T(L.C1_X, L.Y_SEC_TITLE, L.C1_W, [P(B(COL1_TITLE), size=28, lh=1.05)], "Titulo col 1")
    y = L.Y_SEC_TITLE + h1 + 0.1
    T(L.C1_X, y, L.C1_W, COL1, "Corpo col 1", h=L.BOTTOM - y, ajustavel=True)

    # coluna 2: diagrama, tabela, figura do PP1 e, à direita, vazamento e app
    T(L.C2_X, L.Y_SEC_TITLE, L.C2_W, [P(B(COL2_TITLE), size=28, lh=1.05)], "Titulo col 2")
    I("fig_poster_metodo.png", L.C2_X, L.DIAG_Y, L.DIAG_W)
    lh_tab = 20 * 1.1 * PT_CM
    xs = [L.C2_X + sum(L.TAB_COLS[:j]) for j in range(len(L.TAB_COLS))]
    for i, row in enumerate(TABLE):
        ry = L.TAB_Y + i * L.TAB_ROW_H
        for j, cel in enumerate(row):
            T(xs[j] + 0.2, ry + (L.TAB_ROW_H - lh_tab) / 2, L.TAB_COLS[j] - 0.4, [P(cel)], f"Tabela {i + 1}.{j + 1}",
              h=lh_tab + 0.01)    # cada célula em uma linha só
    for i in range(len(TABLE) + 1):     # linhas da tabela: grossas em cima e embaixo, finas entre as bases
        grossa = i in (0, len(TABLE))
        esp = 0.05 if grossa else 0.025
        el.append(("rule", L.C2_X, L.TAB_Y + i * L.TAB_ROW_H - esp / 2, L.C2_W, esp,
                   INK if grossa else "BFBFBF", {}))
    I("fig_poster_exp5b.png", L.C2_X, L.F5B_Y, L.F5B_W)
    y = L.LOW_Y
    y += T(L.R_X, y, L.R_W, LEAK, "Vazamento") + 0.2
    reserva = APP_RESERVA_LINHAS * 20 * 1.1 * PT_CM
    T(L.R_X, y, L.R_W, APP, "App", h=L.BOTTOM - y - reserva, ajustavel=True)

    # coluna 3: frase-síntese, conclusões e, no pé, as referências
    T(L.C3_X, L.Y_SEC_TITLE, L.C3_W, [P(B(COL3_TITLE), size=28, lh=1.05)], "Titulo col 3")
    hr = altura_cm(preparar(REFS, L.C3_W))
    y_ref = L.BOTTOM - hr
    y = L.Y_BODY
    y += T(L.C3_X, y, L.C3_W, SINTESE, "Sintese") + 0.6
    T(L.C3_X, y, L.C3_W, COL3, "Corpo col 3", h=y_ref - 0.2 - y, ajustavel=True)
    T(L.C3_X, y_ref, L.C3_W, REFS, "Referencias", cls="ref")

    # checagens de geometria: tudo dentro da página e do fundo das colunas
    for kind, x, y, w, h, _p, kw in el:
        assert x >= 0 and x + w <= SLIDE_W + 1e-6 and y >= 0 and y + h <= L.COL_BOTTOM + 1e-6, (kind, kw, x, y, w, h)
    return el


# ======================================================================================
# PPTX
# ======================================================================================
KEEP_IDS = {84, 85, 86, 115, 116}   # fundo em degradê (3 partes) e os dois logos


def clear_template(slide):
    tree = slide.shapes._spTree
    for sh in list(slide.shapes):
        if sh.shape_id not in KEEP_IDS:
            tree.remove(sh._element)
    # Remove também as RELAÇÕES das imagens apagadas: sem isto, as figuras do pôster da Entrega 01
    # (com quadros do NTHU-DDD e da UTA-RLDD, isto é, ROSTOS) continuariam embutidas em ppt/media do
    # .pptx, invisíveis mas extraíveis. Só cai a relação de imagem que nenhum elemento restante usa.
    xml = etree.tostring(slide._element).decode("utf-8")
    part = slide.part
    for rId, rel in list(part.rels.items()):
        if rel.reltype.endswith("/image") and f'"{rId}"' not in xml:
            part.rels.pop(rId)


def _sem_estilo(shape, cor):
    """Retângulo/caixa sem o estilo do tema (que pode trazer sombra) e sem contorno."""
    sp = shape._element
    st = sp.find(qn("p:style"))
    if st is not None:
        sp.remove(st)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(cor)
    shape.line.fill.background()
    shape.shadow.inherit = False


def add_card(slide, x, y, w, h, depois_de):
    """Caixa branca de cantos arredondados (raio do template). Os grupos de forma livre do template não
    aparecem no Quick Look do macOS; um retângulo arredondado simples aparece em todos os programas."""
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    sh.adjustments[0] = L.CARD_R / min(w, h)
    sh.name = "Caixa branca"
    _sem_estilo(sh, "FFFFFF")
    depois_de.addnext(sh._element)     # logo acima do fundo, abaixo dos logos e do conteúdo
    return sh


def _set_run(run, text, bold, size, color):
    run.text = text
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "pt-BR")
    for tag in ("a:ea", "a:cs", "a:sym"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def fill_paragraph(par, spec):
    par.alignment = PP_ALIGN.CENTER if spec["align"] == "c" else PP_ALIGN.LEFT
    if spec["sb"]:
        par.space_before = Pt(spec["sb"])
    # espaçamento EXATO em pt (= line-height do HTML): as duas saídas ocupam a mesma altura
    par.line_spacing = Pt(linha_pt(spec))
    etree.SubElement(par._p.get_or_add_pPr(), qn("a:buNone"))
    for i, linha in enumerate(spec["linhas"]):
        if i:
            br = par._p.add_br()
            rPr = br.get_or_add_rPr()
            rPr.set("lang", "pt-BR")
            rPr.set("sz", str(int(spec["size"] * 100)))
        for text, bold, color, size in linha:
            _set_run(par.add_run(), text, bold, size, color)


def add_textbox(slide, x, y, w, h, paragraphs, name):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, spec in enumerate(paragraphs):
        fill_paragraph(tf.paragraphs[0] if i == 0 else tf.add_paragraph(), spec)
    return tb


def add_rule(slide, x, y, w, h, cor):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    sh.name = "Linha da tabela"
    _sem_estilo(sh, cor)
    return sh


def build_pptx(el):
    prs = Presentation(TEMPLATE)
    assert abs(prs.slide_width / 360000 - SLIDE_W) < 0.01
    slide = prs.slides[0]
    clear_template(slide)
    fundo = [sh._element for sh in slide.shapes if sh.shape_id in (84, 85, 86)][-1]
    for kind, x, y, w, h, payload, kw in el:
        if kind == "card":
            fundo = add_card(slide, x, y, w, h, fundo)._element
        elif kind == "text":
            add_textbox(slide, x, y, w, h, payload, kw["name"])
        elif kind == "img":
            pic = slide.shapes.add_picture(os.path.join(ASSETS, payload), Cm(x), Cm(y), Cm(w), Cm(h))
            pic.name = payload
        elif kind == "rule":
            add_rule(slide, x, y, w, h, payload)
    # metadados coerentes (o template trazia autor e data de 05/09)
    agora = datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None, microsecond=0)
    cp = prs.core_properties
    cp.title = TITULO
    cp.author = "; ".join(AUTORES_COMPLETOS)
    cp.last_modified_by = AUTORES_COMPLETOS[0]
    cp.subject = ASSUNTO
    cp.keywords = PALAVRAS
    cp.language = "pt-BR"
    cp.category = "Pôster"
    cp.comments = ""
    cp.created = agora
    cp.modified = agora
    cp.revision = 1
    prs.save(OUT_PPTX)
    return OUT_PPTX


# ======================================================================================
# HTML (mesma geometria, em cm)
# ======================================================================================
def paras_html(paragraphs):
    parts = []
    for spec in paragraphs:
        style = (f"font-size:{spec['size']}pt;line-height:{linha_pt(spec)}pt;color:#{spec['color']};"
                 f"margin-top:{spec['sb']}pt;text-align:{'center' if spec['align'] == 'c' else 'left'};")
        linhas = []
        for ln in spec["linhas"]:
            out = []
            for text, bold, color, size in ln:
                t = html.escape(text).replace(NB, "&nbsp;")
                st = [f"color:#{color}"] + ([f"font-size:{size}pt"] if size != spec["size"] else [])
                t = f'<span style="{";".join(st)}">{t}</span>'
                out.append(f"<b>{t}</b>" if bold else t)
            linhas.append("".join(out))
        parts.append(f'<p style="{style}">' + "<br>".join(linhas) + "</p>")
    return "\n".join(parts)


def box(x, y, w, h, inner, cls="txt", extra=""):
    return (f'<div class="{cls}" style="left:{x}cm;top:{y}cm;width:{w}cm;height:{h}cm;{extra}">'
            f"{inner}</div>")


def build_html(el):
    parts = []
    for kind, x, y, w, h, payload, kw in el:
        if kind == "card":
            parts.append(box(x, y, w, h, "", cls="card"))
    parts.append('<img class="abs" src="assets/logo_pucpr.png" style="left:1.02cm;top:1.70cm;width:6.77cm;height:3.06cm">')
    parts.append('<img class="abs" src="assets/logo_bcc.png" style="left:8.32cm;top:0.96cm;width:4.53cm;height:4.51cm">')
    for kind, x, y, w, h, payload, kw in el:
        if kind == "text":
            parts.append(box(x, y, w, h, paras_html(payload), cls="txt"))
        elif kind == "img":
            parts.append(f'<img class="abs" src="assets/{payload}" style="left:{x}cm;top:{y}cm;width:{w}cm;height:{h}cm">')
        elif kind == "rule":
            parts.append(box(x, y, w, h, "", cls="abs", extra=f"background:#{payload};"))
    autores = html.escape(", ".join(AUTORES_COMPLETOS))
    doc = f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<title>{html.escape(TITULO)}</title>
<meta name="author" content="{autores}">
<meta name="description" content="{html.escape(ASSUNTO)}">
<style>
  @page {{ size: {SLIDE_W}cm {SLIDE_H}cm; margin: 0; }}
  html, body {{ margin: 0; padding: 0; background: #FFFFFF; }}
  body {{ font-family: {FONT}, "Helvetica Neue", Helvetica, sans-serif; color: #{INK}; }}
  .page {{ position: relative; width: {SLIDE_W}cm; height: {SLIDE_H}cm; overflow: hidden;
          background: #9654FF url("assets/fundo_gradiente.png") no-repeat; background-size: 100% 100%;
          page-break-after: avoid; }}
  .abs, .card, .txt {{ position: absolute; box-sizing: border-box; }}
  .card {{ background: #FFFFFF; border-radius: {L.CARD_R}cm; }}
  .txt {{ overflow: visible; }}
  p {{ margin: 0; white-space: nowrap; }}
  img {{ display: block; }}
</style>
</head>
<body>
<div class="page">
{chr(10).join(parts)}
</div>
</body>
</html>
"""
    with open(OUT_HTML, "w", encoding="utf-8") as f:
        f.write(doc)
    return OUT_HTML


# ======================================================================================
# PDF (Chrome headless, mudo), metadados e conferência
# ======================================================================================
def imprimir_pdf():
    if not os.path.exists(CHROME):
        raise SystemExit(f"Chrome não encontrado em {CHROME} (defina CHROME=...)")
    perfil = tempfile.mkdtemp(prefix="perfil_chrome_")
    if os.path.exists(OUT_PDF):
        os.remove(OUT_PDF)
    cmd = [CHROME, "--headless", "--disable-gpu", "--mute-audio", "--no-first-run", "--no-default-browser-check",
           f"--user-data-dir={perfil}", "--no-pdf-header-footer", f"--print-to-pdf={OUT_PDF}",
           "file://" + OUT_HTML]
    proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    # O Chrome headless às vezes grava o PDF e não termina: espera o arquivo ficar estável (até 120 s).
    anterior = -1
    for _ in range(120):
        time.sleep(1)
        if os.path.exists(OUT_PDF) and os.path.getsize(OUT_PDF) > 0:
            atual = os.path.getsize(OUT_PDF)
            if atual == anterior:
                break
            anterior = atual
        if proc.poll() is not None and os.path.exists(OUT_PDF):
            break
    proc.kill()
    proc.wait()
    subprocess.run(["rm", "-rf", perfil], check=False)
    if not os.path.exists(OUT_PDF) or os.path.getsize(OUT_PDF) == 0:
        raise SystemExit("o Chrome headless não gerou poster/poster_final.pdf")

    import pymupdf
    doc = pymupdf.open(OUT_PDF)
    meta = dict(doc.metadata)
    meta.update(title=TITULO, author="; ".join(AUTORES_COMPLETOS), subject=ASSUNTO, keywords=PALAVRAS)
    doc.set_metadata(meta)
    tmp = OUT_PDF + ".tmp"
    doc.save(tmp, garbage=3, deflate=True)
    doc.close()
    os.replace(tmp, OUT_PDF)
    doc = pymupdf.open(OUT_PDF)
    p = doc[0]
    print(f"poster_final.pdf: {len(doc)} página(s), {p.rect.width / 72 * 2.54:.2f} x {p.rect.height / 72 * 2.54:.2f} cm")
    p.get_pixmap(dpi=60).save(OUT_PNG)
    print("poster_final_preview.png atualizado")


TRACOS = (chr(0x2014), chr(0x2013))   # travessão e meia-risca: proibidos no texto visível


def conferir_saidas(el, com_pdf):
    """Fonte mínima (20 pt; 16 pt nas referências e no endereço), nenhum travessão, nenhuma linha
    quebrada de outro jeito no PDF e metadados do PDF com título e autores."""
    erros = []
    esperadas = []
    for kind, x, y, w, h, payload, kw in el:
        if kind != "text":
            continue
        for sp in payload:
            for ln in sp["linhas"]:
                txt = "".join(t for t, *_ in ln)
                esperadas.append(txt.replace(NB, " ").strip())
                for t, n, c, s in ln:
                    minimo = 16 if kw.get("cls") == "ref" else 20
                    if s < minimo:
                        erros.append(f"{kw['name']}: {t!r} com {s} pt")
                    if any(tr in t for tr in TRACOS):
                        erros.append(f"{kw['name']}: travessão em {t!r}")
    # .pptx: tamanhos e traços lidos do próprio arquivo
    prs = Presentation(OUT_PPTX)
    menor_pptx, menor_pptx_corpo = 99, 99
    for sh in prs.slides[0].shapes:
        if sh.shape_id not in KEEP_IDS and (sh.left < 0 or sh.top < 0 or sh.left + sh.width > prs.slide_width
                                            or sh.top + sh.height > prs.slide_height):
            erros.append(f"pptx: {sh.name} fora da página")
        if not sh.has_text_frame:
            continue
        for par in sh.text_frame.paragraphs:
            for r in par.runs:
                if not r.text.strip():
                    continue
                pt = r.font.size.pt
                menor_pptx = min(menor_pptx, pt)
                if sh.name not in ("Referencias", "Endereco do repositorio"):
                    menor_pptx_corpo = min(menor_pptx_corpo, pt)
                if any(tr in r.text for tr in TRACOS):
                    erros.append(f"pptx: travessão em {r.text!r}")
    print(f"pptx: menor fonte {menor_pptx:g} pt (fora referências e endereço: {menor_pptx_corpo:g} pt)")
    if menor_pptx_corpo < 20 or menor_pptx < 16:
        erros.append("pptx: fonte abaixo do mínimo")
    if com_pdf:
        import pymupdf
        doc = pymupdf.open(OUT_PDF)
        pg = doc[0]
        linhas_pdf, menor, menor_corpo = [], 99, 99
        for b in pg.get_text("dict")["blocks"]:
            for ln in b.get("lines", []):
                txt = "".join(s["text"] for s in ln["spans"]).replace(chr(0xA0), " ").strip()
                if txt:
                    linhas_pdf.append(txt)
                for s in ln["spans"]:
                    if not s["text"].strip():
                        continue
                    menor = min(menor, s["size"])
                    if s["size"] < 19.5 and not (s["text"].startswith("[") or "Código e app" in txt or
                                                 REPO_URL in txt or any(r in txt for r in ("et al.", "ABRAMET", "Fusek"))):
                        erros.append(f"pdf: {s['text']!r} com {s['size']:.1f} pt")
                    menor_corpo = min(menor_corpo, s["size"]) if s["size"] >= 19.5 else menor_corpo
        texto = pg.get_text()
        for tr in TRACOS:
            if tr in texto:
                erros.append(f"pdf: {texto.count(tr)} ocorrência(s) de {tr!r}")
        faltam = [e for e in esperadas if e and e not in linhas_pdf]
        if faltam:
            erros.append(f"pdf: {len(faltam)} linha(s) quebradas de outro jeito, ex.: {faltam[:3]}")
        md = doc.metadata
        if md.get("title") != TITULO or not md.get("author", "").startswith(AUTORES_COMPLETOS[0]):
            erros.append("pdf: metadados de título/autor ausentes")
        print(f"pdf: menor fonte {menor:.1f} pt; {len(esperadas)} linhas de texto, todas iguais às do .pptx"
              if not faltam else f"pdf: menor fonte {menor:.1f} pt")
    if erros:
        raise SystemExit("conferência falhou:\n  " + "\n  ".join(erros))
    print("conferência: ok (fontes, traços, quebras de linha e metadados)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--pdf", action="store_true", help="imprime o PDF pelo Chrome headless e confere")
    args = ap.parse_args()
    els = elements()
    print("Samsung medido:", SAMSUNG_MEDIDO)
    print(build_pptx(els))
    print(build_html(els))
    if args.pdf:
        imprimir_pdf()
    conferir_saidas(els, args.pdf)
